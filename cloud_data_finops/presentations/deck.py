"""Deterministic executive deck built with python-pptx.

The deck uses only native shapes, text, and tables (no embedded images or
charts), so its bytes depend only on the findings, the telemetry, and the
pinned python-pptx version. Core properties and ZIP timestamps are fixed.
"""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
GREEN = RGBColor(0x10, 0xB9, 0x81)
NEUTRAL = RGBColor(0xE2, 0xE8, 0xF0)
SLATE = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xF8, 0xFA, 0xFC)
FONT = "Calibri"
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_WIDTH = SLIDE_WIDTH - 2 * MARGIN
ROWS_PER_TABLE_SLIDE = 11
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
TITLE_ONLY_LAYOUT = 5

# Section names are part of the deck contract checked by tests.
SECTIONS = (
    "title",
    "executive-summary",
    "cost-drivers",
    "prioritized-recommendations",
    "risks-and-assumptions",
    "evidence-appendix",
)


def _style_run(run: Any, size: float, *, bold: bool = False, color: RGBColor = NAVY) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _text_box(slide: Any, left: int, top: int, width: int, height: int, paragraphs: list[tuple[str, float, bool, RGBColor]], *, name: str) -> Any:
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.name = name
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = Inches(0.05)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    for index, (text, size, bold, color) in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(size * 0.45)
        run = paragraph.add_run()
        run.text = text
        _style_run(run, size, bold=bold, color=color)
    return shape


def _rectangle(slide: Any, left: int, top: int, width: int, height: int, color: RGBColor, *, name: str, description: str = "") -> Any:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, max(int(width), 1), height)
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    if description:
        shape._element.nvSpPr.cNvPr.set("descr", description)
    return shape


def _new_slide(presentation: Any, title: str, section: str, *, dark: bool = False) -> Any:
    slide = presentation.slides.add_slide(presentation.slide_layouts[TITLE_ONLY_LAYOUT])
    slide.name = section
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY if dark else WHITE
    title_shape = slide.shapes.title
    title_shape.left, title_shape.top = MARGIN, Inches(0.35) if not dark else Inches(2.2)
    title_shape.width, title_shape.height = CONTENT_WIDTH, Inches(0.8) if not dark else Inches(1.2)
    frame = title_shape.text_frame
    frame.margin_left = Inches(0.05)
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.text = title
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    _style_run(paragraph.runs[0], 40 if dark else 28, bold=True, color=WHITE if dark else NAVY)
    accent_top = Inches(3.5) if dark else Inches(1.18)
    _rectangle(slide, MARGIN, accent_top, Inches(1.2), Inches(0.06), CYAN, name="Accent")
    return slide


def _footer(slide: Any, specification: dict[str, Any], page: int) -> None:
    text = f"Synthetic fixture data · {specification['assessment_id']} · {page}"
    _text_box(slide, MARGIN, Inches(7.0), CONTENT_WIDTH, Inches(0.35), [(text, 10, False, SLATE)], name="Footer")


def _table(slide: Any, rows: list[list[str]], widths: list[float], top: int, *, font_size: float = 12, row_height: float = 0.42) -> Any:
    shape = slide.shapes.add_table(len(rows), len(widths), MARGIN, top, Inches(sum(widths)), Inches(row_height * len(rows)))
    shape.name = "Table"
    table = shape.table
    for column, width in enumerate(widths):
        table.columns[column].width = Inches(width)
    for row_index, values in enumerate(rows):
        table.rows[row_index].height = Inches(row_height)
        for column, value in enumerate(values):
            cell = table.cell(row_index, column)
            cell.fill.solid()
            header = row_index == 0
            cell.fill.fore_color.rgb = NAVY if header else (PALE if row_index % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame = cell.text_frame
            frame.word_wrap = True
            frame.text = value
            _style_run(frame.paragraphs[0].runs[0], font_size, bold=header, color=WHITE if header else NAVY)
    return shape


def _chunks(items: list[Any], size: int) -> list[list[Any]]:
    return [items[index : index + size] for index in range(0, len(items), size)] or [[]]


def headline_metric(finding: dict[str, Any]) -> str:
    """One short, unit-bearing line per finding for the appendix table."""
    result = finding["calculation"]["result"]
    evidence = finding["observed_evidence"]
    rule = finding["rule_id"]
    if rule in {"BQ-001", "BQ-003"}:
        return f"{result['terabytes_billed']:,.2f} TB billed"
    if rule == "BQ-002":
        return f"{evidence['run_count']} runs, {result['repeated_terabytes']:,.1f} TB repeated"
    if rule == "BQ-004":
        return f"{evidence['runs_per_day']:g} runs/day; source changes {evidence['source_updates_per_day']:g}/day"
    if rule == "BQ-005":
        return f"{result['utilization']:.0%} of {evidence['baseline_slots']} slots used"
    if rule in {"AWS-001", "AWS-002"}:
        return f"USD {evidence['cost_usd']:,.0f} vs USD {evidence['baseline_usd']:,.0f} baseline"
    if rule == "AWS-003":
        return f"{evidence['utilization']:.0%} utilization over {evidence['observed_days']} days"
    if rule == "CMT-001":
        return f"CV {result['coefficient_of_variation']:.2f}, coverage {evidence['coverage_ratio']:.0%}"
    return "see findings.json"


def estimate_label(finding: dict[str, Any]) -> str:
    impact = finding["estimated_impact"]
    if impact["value"] is None:
        return "Not estimated"
    value = impact["value"]
    labels = {
        "BQ-002": f"{value:,.1f} TB (upper bound)",
        "BQ-004": f"{value:,.1f} TB per 30 days",
        "BQ-005": f"{value:,.0f} idle slot-hours",
        "AWS-003": f"USD {value:,.0f} per month",
        "CMT-001": f"USD {value:,.0f} per month",
    }
    return labels.get(finding["rule_id"], f"{value:,.1f}")


def executive_messages(telemetry: dict[str, Any], findings: list[dict[str, Any]]) -> list[str]:
    providers = Counter(finding["provider"] for finding in findings)
    messages = []
    bq = [finding for finding in findings if finding["rule_id"] == "BQ-001"]
    if bq:
        top = max(bq, key=lambda finding: finding["calculation"]["result"]["terabytes_billed"])
        pruning = " without partition pruning" if "BQ-003" in top["related_rules"] else ""
        messages.append(f"GCP: {providers['gcp']} findings. The largest job billed {top['calculation']['result']['terabytes_billed']:,.1f} TB{pruning}.")
    aws = [finding for finding in findings if finding["rule_id"] == "AWS-001"]
    if aws:
        top = max(aws, key=lambda finding: finding["calculation"]["result"]["increase_usd"])
        messages.append(f"AWS: {providers['aws']} findings. {top['subject']} rose USD {top['calculation']['result']['increase_usd']:,.0f} above its baseline.")
    estimated = sum(1 for finding in findings if finding["estimated_impact"]["value"] is not None)
    messages.append(f"{estimated} of {len(findings)} findings carry an upper-bound estimate in TB, slot-hours, or USD; none is a saving.")
    if any(finding["rule_id"] == "CMT-001" and finding["estimated_impact"]["value"] is None for finding in findings):
        messages.append("No commitment purchase is recommended: the commitment evidence is incomplete.")
    return messages


def _title_slide(presentation: Any, telemetry: dict[str, Any]) -> None:
    slide = _new_slide(presentation, "Synthetic cloud data FinOps assessment", "title", dark=True)
    period = telemetry["period"]
    _text_box(
        slide,
        MARGIN,
        Inches(3.75),
        CONTENT_WIDTH,
        Inches(1.6),
        [
            (f"Assessment {telemetry['assessment_id']} · {period['start']} to {period['end']} · GCP and AWS", 20, False, WHITE),
            ("Synthetic fixture data. Nothing here describes a real environment, bill, or saving.", 16, True, CYAN),
        ],
        name="Subtitle",
    )


def _summary_slide(presentation: Any, telemetry: dict[str, Any], findings: list[dict[str, Any]]) -> None:
    slide = _new_slide(presentation, "Executive summary", "executive-summary")
    priorities = Counter(finding["priority"] for finding in findings)
    tiles = [("Findings", len(findings), BLUE), ("High", priorities["high"], NAVY), ("Medium", priorities["medium"], BLUE), ("Low", priorities["low"], GREEN)]
    for index, (label, value, color) in enumerate(tiles):
        top = Inches(1.55 + index * 1.3)
        _rectangle(slide, MARGIN, top, Inches(3.2), Inches(1.1), PALE, name=f"Tile {label}", description=f"{label}: {value}")
        _rectangle(slide, MARGIN, top, Inches(0.09), Inches(1.1), color, name=f"Tile marker {label}")
        _text_box(
            slide,
            MARGIN + Inches(0.3),
            top + Inches(0.04),
            Inches(2.8),
            Inches(1.02),
            [(str(value), 28, True, NAVY), (f"{label} priority" if label != "Findings" else "Findings in total", 12, False, SLATE)],
            name=f"Tile text {label}",
        )
    paragraphs = [(message, 18, False, NAVY) for message in executive_messages(telemetry, findings)]
    _text_box(slide, MARGIN + Inches(3.7), Inches(1.55), Inches(8.4), Inches(5.0), paragraphs, name="Key messages")


def _bar_rows(slide: Any, left: int, top: int, width: int, title: str, rows: list[tuple[str, float, float | None, str]], section: str) -> None:
    """Draw labelled horizontal bars. Each row: label, value, optional baseline, value text."""
    _text_box(slide, left, top, width, Inches(0.5), [(title, 16, True, NAVY)], name=f"{section} heading")
    scale = max([max(value, baseline or 0) for _, value, baseline, _ in rows] or [1]) or 1
    label_width = Inches(2.3)
    bar_left = left + label_width
    bar_space = width - label_width - Inches(1.6)
    for index, (label, value, baseline, value_text) in enumerate(rows):
        row_top = top + Inches(0.7 + index * 0.85)
        _text_box(slide, left, row_top, label_width - Inches(0.1), Inches(0.55), [(label, 12, False, NAVY)], name=f"{section} label {index + 1}")
        if baseline is not None:
            _rectangle(
                slide,
                bar_left,
                row_top + Inches(0.05),
                int(bar_space * baseline / scale),
                Inches(0.18),
                NEUTRAL,
                name=f"{section} baseline {index + 1}",
                description=f"{label} baseline {baseline:,.0f}",
            )
            _rectangle(
                slide,
                bar_left,
                row_top + Inches(0.27),
                int(bar_space * value / scale),
                Inches(0.18),
                BLUE,
                name=f"{section} bar {index + 1}",
                description=f"{label} {value_text}",
            )
        else:
            _rectangle(
                slide,
                bar_left,
                row_top + Inches(0.1),
                int(bar_space * value / scale),
                Inches(0.3),
                CYAN,
                name=f"{section} bar {index + 1}",
                description=f"{label} {value_text}",
            )
        _text_box(
            slide, bar_left + bar_space + Inches(0.1), row_top, Inches(1.5), Inches(0.55), [(value_text, 12, True, NAVY)], name=f"{section} value {index + 1}"
        )


def _drivers_slide(presentation: Any, telemetry: dict[str, Any]) -> None:
    slide = _new_slide(presentation, "Cost drivers by provider", "cost-drivers")
    jobs = sorted(telemetry["gcp"].get("jobs", []), key=lambda job: (-job["bytes_billed"], job["job_id"]))[:5]
    job_rows = [(job["job_id"], job["bytes_billed"] / 10**12, None, f"{job['bytes_billed'] / 10**12:,.2f} TB") for job in jobs]
    costs = sorted(telemetry["aws"].get("costs", []), key=lambda row: (-row["cost_usd"], row["service"]))[:5]
    cost_rows = [(row["service"], row["cost_usd"], row["baseline_usd"], f"USD {row['cost_usd']:,.0f}") for row in costs]
    half = (CONTENT_WIDTH - Inches(0.5)) // 2
    _bar_rows(slide, MARGIN, Inches(1.45), half, "GCP: bytes billed, five largest jobs (TB)", job_rows, "GCP")
    _bar_rows(slide, MARGIN + half + Inches(0.5), Inches(1.45), half, "AWS: cost vs baseline by service (USD)", cost_rows, "AWS")
    _text_box(
        slide,
        MARGIN,
        Inches(6.35),
        CONTENT_WIDTH,
        Inches(0.55),
        [("Units differ by panel and are not compared. Grey bars show the declared baseline; blue bars show the assessment period.", 12, False, SLATE)],
        name="Drivers note",
    )


def _recommendation_slides(presentation: Any, findings: list[dict[str, Any]]) -> None:
    chunks = _chunks(findings, ROWS_PER_TABLE_SLIDE)
    for index, chunk in enumerate(chunks):
        suffix = f" ({index + 1} of {len(chunks)})" if len(chunks) > 1 else ""
        slide = _new_slide(presentation, f"Prioritized recommendations{suffix}", "prioritized-recommendations")
        rows = [["#", "Priority", "Rule", "Subject", "Next action", "Confidence"]]
        offset = index * ROWS_PER_TABLE_SLIDE
        for position, finding in enumerate(chunk, start=offset + 1):
            rows.append([str(position), finding["priority"], finding["rule_id"], finding["subject"], finding["action"], finding["confidence"]])
        _table(slide, rows, [0.5, 1.0, 1.0, 3.0, 5.2, 1.4], Inches(1.45))


def _risks_slide(presentation: Any) -> None:
    slide = _new_slide(presentation, "Risks and assumptions", "risks-and-assumptions")
    assumptions = [
        "Telemetry is synthetic and generated from a fixed seed.",
        "No pricing model is applied; estimates are upper bounds in TB, slot-hours, or USD.",
        "Thresholds are catalog defaults and need review for a real workload.",
        "Resource utilization and commitment data are optional telemetry and were approved in the fixture scope.",
    ]
    risks = [
        "Acting on an upper bound can overstate the benefit of a change.",
        "Reducing reservation baselines can slow peak workloads; review peak slot usage first.",
        "Retiring an underutilized resource can break an undocumented dependency.",
        "Buying a commitment on incomplete history can lock in unused spend.",
    ]
    half = (CONTENT_WIDTH - Inches(0.5)) // 2
    for index, (heading, items, color) in enumerate((("Assumptions", assumptions, BLUE), ("Risks before acting", risks, GREEN))):
        left = MARGIN + index * (half + Inches(0.5))
        _rectangle(slide, left, Inches(1.5), half, Inches(4.0), PALE, name=f"{heading} panel")
        _rectangle(slide, left, Inches(1.5), half, Inches(0.08), color, name=f"{heading} marker")
        paragraphs = [(heading, 18, True, NAVY)] + [(f"• {item}", 15, False, NAVY) for item in items]
        _text_box(slide, left + Inches(0.25), Inches(1.75), half - Inches(0.5), Inches(3.6), paragraphs, name=f"{heading} text")


def _appendix_slides(presentation: Any, specification: dict[str, Any], findings: list[dict[str, Any]], findings_sha256: str) -> None:
    chunks = _chunks(findings, ROWS_PER_TABLE_SLIDE)
    for index, chunk in enumerate(chunks):
        suffix = f" ({index + 1} of {len(chunks) + 1})"
        slide = _new_slide(presentation, f"Evidence appendix{suffix}", "evidence-appendix")
        rows = [["Rule", "Subject", "Evidence source", "Observed metric", "Estimate"]]
        rows.extend(
            [finding["rule_id"], finding["subject"], finding["evidence_source"], headline_metric(finding), estimate_label(finding)] for finding in chunk
        )
        _table(slide, rows, [1.0, 3.0, 1.9, 3.7, 2.5], Inches(1.45))
    slide = _new_slide(presentation, f"Evidence appendix ({len(chunks) + 1} of {len(chunks) + 1})", "evidence-appendix")
    lines = [
        ("Every number in this deck comes from findings.json and the synthetic fixture.", 16, False, NAVY),
        (f"findings.json SHA-256: {findings_sha256}", 13, True, NAVY),
        (f"Generator: cloud_data_finops.synthetic, seed {specification['scope'].get('generator', {}).get('seed', 'n/a')}", 14, False, NAVY),
        ("Reproduce: make reproduce (see docs/reproduce.md for expected checksums)", 14, False, NAVY),
        ("Full per-finding evidence, formulas, units, and assumptions: report.md", 14, False, NAVY),
    ]
    _text_box(slide, MARGIN, Inches(1.55), CONTENT_WIDTH, Inches(4.8), lines, name="Reproduction")


def findings_digest(findings: list[dict[str, Any]]) -> str:
    """SHA-256 of findings.json exactly as report.write_findings serializes it."""
    return hashlib.sha256((json.dumps(findings, indent=2, sort_keys=True) + "\n").encode("utf-8")).hexdigest()


def _fixed_timestamp(specification: dict[str, Any]) -> datetime:
    return datetime.strptime(specification["period"]["end"], "%Y-%m-%d")


def build_presentation(telemetry: dict[str, Any], findings: list[dict[str, Any]]) -> Any:
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    _title_slide(presentation, telemetry)
    _summary_slide(presentation, telemetry, findings)
    _drivers_slide(presentation, telemetry)
    _recommendation_slides(presentation, findings)
    _risks_slide(presentation)
    _appendix_slides(presentation, telemetry, findings, findings_digest(findings))
    for page, slide in enumerate(presentation.slides, start=1):
        if page > 1:
            _footer(slide, telemetry, page)

    properties = presentation.core_properties
    timestamp = _fixed_timestamp(telemetry)
    properties.title = "Synthetic cloud data FinOps assessment"
    properties.subject = f"Assessment {telemetry['assessment_id']} (synthetic fixture data)"
    properties.author = "cloud-data-finops-sdd-toolkit"
    properties.last_modified_by = "cloud-data-finops-sdd-toolkit"
    properties.keywords = "FinOps; GCP; AWS; synthetic"
    properties.comments = "Generated from findings.json by cloud_data_finops.presentations."
    properties.created = timestamp
    properties.modified = timestamp
    properties.last_printed = timestamp
    properties.revision = 1
    return presentation


def normalize_package(raw: bytes) -> bytes:
    """Rewrite the OPC ZIP with fixed timestamps, attributes, and no compression so equal content gives equal bytes."""
    source = zipfile.ZipFile(io.BytesIO(raw))
    target_buffer = io.BytesIO()
    with zipfile.ZipFile(target_buffer, "w") as target:
        for item in source.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=FIXED_ZIP_TIME)
            # Stored (uncompressed) entries avoid differences between zlib builds.
            info.compress_type = zipfile.ZIP_STORED
            info.external_attr = 0o600 << 16
            info.create_system = 0
            target.writestr(info, source.read(item.filename))
    return target_buffer.getvalue()


def render_deck(telemetry: dict[str, Any], findings: list[dict[str, Any]]) -> bytes:
    buffer = io.BytesIO()
    build_presentation(telemetry, findings).save(buffer)
    return normalize_package(buffer.getvalue())


def write_deck(path: Path, telemetry: dict[str, Any], findings: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(render_deck(telemetry, findings))


def slide_outline(deck: bytes | Path) -> list[dict[str, Any]]:
    """Titles, sections, and visible text per slide, in reading order. Used for golden snapshots."""
    source = io.BytesIO(deck) if isinstance(deck, bytes) else deck
    presentation = Presentation(source)
    outline = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(cell.text for cell in row.cells))
        outline.append({"index": index, "section": slide.name, "title": slide.shapes.title.text, "text": texts})
    return outline
