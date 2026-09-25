"""Presentation tests: the generated deck opens, has the required sections, accessible titles, and text that fits."""

import math
import unittest
import zipfile
from io import BytesIO

from pptx import Presentation
from pptx.util import Pt

from cloud_data_finops.policies import evaluate
from cloud_data_finops.presentations import SECTIONS, render_deck
from cloud_data_finops.workflow import collect_fixture_telemetry
from tests.helpers import ROOT, load_fixture

EVIDENCE_DECK = ROOT / "docs" / "evidence" / "release-fixture" / "synthetic-finops-assessment.pptx"
# Average Calibri glyph width is close to half the font size; 0.55 leaves a safety margin.
AVERAGE_GLYPH_WIDTH = 0.55
LINE_HEIGHT = 1.2
PALETTE = {"0F172A", "2563EB", "22D3EE", "10B981", "E2E8F0"}


def estimated_text_height(text_frame, width_emu: int) -> int:
    height = 0
    usable = width_emu - (text_frame.margin_left or 0) - (text_frame.margin_right or 0)
    for paragraph in text_frame.paragraphs:
        size = max((run.font.size for run in paragraph.runs if run.font.size), default=Pt(18))
        characters = sum(len(run.text) for run in paragraph.runs)
        lines = max(1, math.ceil(characters * size * AVERAGE_GLYPH_WIDTH / usable))
        height += int(lines * size * LINE_HEIGHT) + int(paragraph.space_after or 0)
    return height + (text_frame.margin_top or 0) + (text_frame.margin_bottom or 0)


class PresentationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        telemetry = collect_fixture_telemetry(load_fixture())
        cls.findings = evaluate(telemetry)
        cls.raw = render_deck(telemetry, cls.findings)
        cls.deck = Presentation(BytesIO(cls.raw))

    def test_deck_opens_with_the_expected_slide_count(self) -> None:
        self.assertEqual(len(self.deck.slides), 7)
        self.assertEqual(len(Presentation(str(EVIDENCE_DECK)).slides), 7)

    def test_required_sections_appear_in_order(self) -> None:
        sections = [slide.name for slide in self.deck.slides]
        self.assertEqual(list(dict.fromkeys(sections)), list(SECTIONS))

    def test_every_slide_has_a_unique_accessible_title(self) -> None:
        titles = []
        for slide in self.deck.slides:
            title = slide.shapes.title
            self.assertIsNotNone(title, slide.name)
            self.assertTrue(title.is_placeholder)
            self.assertTrue(title.text.strip())
            titles.append(title.text)
        self.assertEqual(len(titles), len(set(titles)))
        self.assertEqual(titles[1:5], ["Executive summary", "Cost drivers by provider", "Prioritized recommendations", "Risks and assumptions"])

    def test_every_finding_is_listed_in_recommendations_and_appendix(self) -> None:
        def table_text(section: str) -> str:
            return "\n".join(
                cell.text
                for slide in self.deck.slides
                if slide.name == section
                for shape in slide.shapes
                if shape.has_table
                for row in shape.table.rows
                for cell in row.cells
            )

        recommendations, appendix = table_text("prioritized-recommendations"), table_text("evidence-appendix")
        for finding in self.findings:
            self.assertIn(finding["subject"], recommendations)
            self.assertIn(finding["action"], recommendations)
            self.assertIn(finding["subject"], appendix)

    def test_synthetic_label_is_visible(self) -> None:
        first = "\n".join(shape.text_frame.text for shape in self.deck.slides[0].shapes if shape.has_text_frame)
        self.assertIn("Synthetic fixture data", first)
        for slide in list(self.deck.slides)[1:]:
            footer = [shape for shape in slide.shapes if shape.name == "Footer"]
            self.assertEqual(len(footer), 1, slide.name)
            self.assertTrue(footer[0].text_frame.text.startswith("Synthetic fixture data"))

    def test_uses_the_portfolio_palette_and_no_images_or_charts(self) -> None:
        with zipfile.ZipFile(BytesIO(self.raw)) as package:
            names = package.namelist()
            xml = "".join(package.read(name).decode("utf-8") for name in names if name.startswith("ppt/slides/slide"))
        self.assertFalse([name for name in names if name.startswith(("ppt/media/", "ppt/charts/", "ppt/embeddings/"))])
        for color in PALETTE:
            self.assertIn(f'val="{color}"', xml)

    def test_shapes_stay_inside_the_slide(self) -> None:
        width, height = self.deck.slide_width, self.deck.slide_height
        for slide in self.deck.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0, (slide.name, shape.name))
                self.assertGreaterEqual(shape.top, 0, (slide.name, shape.name))
                self.assertLessEqual(shape.left + shape.width, width, (slide.name, shape.name))
                self.assertLessEqual(shape.top + shape.height, height, (slide.name, shape.name))

    def test_text_fits_its_box(self) -> None:
        for slide in self.deck.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    needed = estimated_text_height(shape.text_frame, shape.width)
                    self.assertLessEqual(needed, shape.height, (slide.name, shape.name, shape.text_frame.text[:60]))
                if shape.has_table:
                    table = shape.table
                    for row_index, row in enumerate(table.rows):
                        for column_index, cell in enumerate(row.cells):
                            frame = cell.text_frame
                            size = frame.paragraphs[0].runs[0].font.size
                            usable = table.columns[column_index].width - cell.margin_left - cell.margin_right
                            self.assertLessEqual(len(cell.text) * size * AVERAGE_GLYPH_WIDTH, usable, (slide.name, row_index, cell.text))

    def test_core_properties_are_fixed(self) -> None:
        properties = self.deck.core_properties
        self.assertEqual(properties.created.isoformat(), "2026-08-31T00:00:00")
        self.assertEqual(properties.modified.isoformat(), "2026-08-31T00:00:00")
        self.assertEqual(properties.revision, 1)
        self.assertIn("synthetic", properties.subject)

    def test_zip_entries_have_fixed_timestamps(self) -> None:
        with zipfile.ZipFile(BytesIO(self.raw)) as package:
            self.assertEqual({info.date_time for info in package.infolist()}, {(1980, 1, 1, 0, 0, 0)})
            self.assertEqual(package.namelist()[0], "[Content_Types].xml")

    def test_many_findings_paginate_tables(self) -> None:
        telemetry = collect_fixture_telemetry(load_fixture())
        findings = self.findings * 2
        deck = Presentation(BytesIO(render_deck(telemetry, findings)))
        titles = [slide.shapes.title.text for slide in deck.slides]
        self.assertIn("Prioritized recommendations (2 of 2)", titles)
        self.assertIn("Evidence appendix (3 of 3)", titles)


if __name__ == "__main__":
    unittest.main()
